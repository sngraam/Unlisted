"""Read Amazon Seller Central XLSM metadata without loading or executing macros.

This extracts a versioned field catalog, not seller product facts or an SP-API schema.
The Template sheet's row 5 is the exact export header; Data Definitions supplies
requirement labels; Dropdown Lists supplies workbook choices. Examples are kept
as source guidance, never category identity: Amazon shows SHIRT for other types.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


PARSER_VERSION = 2
REQUIRED_SHEETS = {"Data Definitions", "Template", "Valid Values", "Browse Data", "Dropdown Lists"}
INDEX_RE = re.compile(r"#\d+")
MARKETPLACE_RE = re.compile(r"marketplace_id=([A-Z0-9]+)")
LANGUAGE_RE = re.compile(r"language_tag=([a-z]{2}_[A-Z]{2})")
REQUIREMENTS = {
    "Required": "REQUIRED",
    "Conditionally Required": "CONDITIONAL",
    "Optional": "OPTIONAL",
    "Recommended": "RECOMMENDED",
}

# Only clear, lossless mappings are shared across category templates. Everything
# else remains addressable by its exact Amazon header and semantic attribute name.
CANONICAL_PATHS = {
    "contribution_sku": "variant.sellerSku",
    "product_type": "listing.productType",
    "recommended_browse_nodes": "listing.browseNodeId",
    "item_name": "revision.title",
    "product_description": "revision.description",
    "bullet_point": "revision.bulletPoints",
    "brand": "product.canonicalData.brand",
    "color": "variant.color",
    "country_of_origin": "variant.countryOfOrigin",
    "fabric_type": "variant.attributes.fabricType",
    "material": "variant.attributes.material",
}


def _text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _pattern(key: str) -> str:
    return INDEX_RE.sub("#*", key)


def _attribute(key: str) -> str:
    return re.split(r"[\[#]", key, maxsplit=1)[0]


def _definitions(sheet) -> tuple[dict[str, dict], list[dict]]:
    if [_text(sheet.cell(2, col).value) for col in range(1, 7)] != [
        "Group Name", "Field Name", "Local Label Name", "Accepted Values", "Example", "Required?"
    ]:
        raise ValueError("Unexpected Data Definitions layout; review parser before importing")
    group = ""
    result = {}
    documentation = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3):
        if row[0] and not row[1]:
            group = _text(row[0])
        if not row[1]:
            continue
        key = _text(row[1])
        requirement = _text(row[5])
        if requirement not in REQUIREMENTS:
            raise ValueError(f"Unknown requirement {requirement!r} for {key}")
        pattern = _pattern(key)
        definition = {
            "label": _text(row[2]),
            "group": group,
            "description": _text(row[3]),
            "requirement": REQUIREMENTS[requirement],
        }
        if pattern in result and result[pattern] != definition:
            raise ValueError(f"Conflicting definition for repeated field {pattern}")
        result[pattern] = definition
        documentation.append({"pattern": pattern, "row": row_number, "example": _text(row[4])})
    return result, documentation


def _suggested_choice_columns(path: Path) -> set[int]:
    """A list without a stopping error is a suggestion, not a closed enum.

    In particular, Amazon puts Delete Offer in the price dropdown while allowing
    numeric prices. Read the workbook's validation flags instead of guessing.
    """
    from openpyxl.utils.cell import range_boundaries
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(path) as archive:
        book = ET.fromstring(archive.read("xl/workbook.xml"))
        relations = {r.attrib["Id"]: r.attrib["Target"] for r in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
        sheet = next(s for s in book.find("m:sheets", ns) if s.attrib["name"] == "Template")
        target = relations[sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]]
        sheet_xml = ET.fromstring(archive.read(target.lstrip("/") if target.startswith("/") else "xl/" + target))
        suggested = set()
        for rule in sheet_xml.findall(".//m:dataValidation", ns):
            if rule.attrib.get("type") != "list":
                continue
            strict = rule.attrib.get("showErrorMessage") in {"1", "true"} and rule.attrib.get("errorStyle", "stop") == "stop"
            if not strict:
                for cell_range in rule.attrib.get("sqref", "").split():
                    first_col, first_row, last_col, last_row = range_boundaries(cell_range)
                    if first_row <= 7 <= last_row:
                        suggested.update(range(first_col, last_col + 1))
        return suggested


def _product_type(sheet) -> str:
    values = set()
    for row in sheet.iter_rows(values_only=True):
        if _text(row[1] if len(row) > 1 else None).startswith("Product Type -"):
            values.update(_text(value) for value in row[2:] if _text(value))
    if len(values) != 1:
        raise ValueError(f"Expected exactly one allowed Product Type; found {sorted(values)}")
    product_type = values.pop()
    if not re.fullmatch(r"[A-Z][A-Z0-9_]{1,119}", product_type):
        raise ValueError(f"Unexpected Product Type format: {product_type!r}")
    return product_type


def _choices(sheet) -> dict[str, list[str]]:
    headers = next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))
    result = {_text(key): [] for index, key in enumerate(headers) if index >= 2 and _text(key)}
    seen = {key: set() for key in result}
    # A/B are workbook support data, not attribute dropdowns.
    for row in sheet.iter_rows(min_row=4, values_only=True):
        for index, raw in enumerate(row):
            if index < 2 or index >= len(headers):
                continue
            key = _text(headers[index])
            if not key:
                continue
            value = _text(raw)
            if value and value not in seen[key]:
                result[key].append(value)
                seen[key].add(value)
    return result


def _browse_nodes(sheet) -> list[dict[str, str]]:
    if [_text(sheet.cell(1, col).value) for col in (1, 2)] != ["Browse Node", "BrowsePath"]:
        raise ValueError("Unexpected Browse Data layout; review parser before importing")
    nodes = []
    seen = set()
    for row in sheet.iter_rows(min_row=2, max_col=2, values_only=True):
        node_id, path = (_text(value) for value in row)
        if node_id and path and node_id not in seen:
            nodes.append({"id": node_id, "path": path})
            seen.add(node_id)
    return nodes


def extract_template(path: str | Path) -> dict:
    path = Path(path)
    if path.suffix.lower() != ".xlsm":
        raise ValueError("Expected an .xlsm Amazon inventory template")
    if path.stat().st_size > 30_000_000:
        raise ValueError("Template exceeds the 30 MB import limit")
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        missing = REQUIRED_SHEETS - set(workbook.sheetnames)
        if missing:
            raise ValueError(f"Missing Amazon template sheets: {sorted(missing)}")
        definitions, documentation = _definitions(workbook["Data Definitions"])
        suggested_columns = _suggested_choice_columns(path)
        product_type = _product_type(workbook["Valid Values"])
        choices = _choices(workbook["Dropdown Lists"])
        sheet = workbook["Template"]
        keys = next(sheet.iter_rows(min_row=5, max_row=5, values_only=True))
        labels = next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))
        if len([key for key in keys if _text(key)]) != len(set(_text(key) for key in keys if _text(key))):
            raise ValueError("Duplicate exact Template headers")
        fields = []
        for index, key_value in enumerate(keys):
            key = _text(key_value)
            if not key:
                continue
            definition = definitions.get(_pattern(key))
            if not definition:
                raise ValueError(f"No Data Definitions entry for Template header {key}")
            attribute = _attribute(key)
            field = {
                "column": index + 1,
                "key": key,
                "pattern": _pattern(key),
                "attribute": attribute,
                "label": _text(labels[index]) or definition["label"],
                **definition,
            }
            if attribute in CANONICAL_PATHS and not (
                attribute == "color" and ".standardized_values" in key
            ):
                field["canonicalPath"] = CANONICAL_PATHS[attribute]
            if key in choices:
                field["allowedValues"] = choices[key]
            if attribute == "product_type":
                field["allowedValues"] = [product_type]
            fields.append(field)
        if len(fields) < 20:
            raise ValueError("Too few Amazon template fields; likely wrong workbook layout")
        marketplace_ids = {match for field in fields for match in MARKETPLACE_RE.findall(field["key"])}
        languages = {match for field in fields for match in LANGUAGE_RE.findall(field["key"])}
        if len(marketplace_ids) != 1 or len(languages) != 1:
            raise ValueError("Template must have exactly one marketplace ID and one language")
        nodes = _browse_nodes(workbook["Browse Data"])
        semantic = {
            "parserVersion": PARSER_VERSION,
            "platform": "AMAZON",
            "marketplaceId": marketplace_ids.pop(),
            "productType": product_type,
            "language": languages.pop(),
            "fields": fields,
            "browseNodes": nodes,
        }
        schema_hash = hashlib.sha256(json.dumps(semantic, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()
        warnings = []
        if path.stem.upper() != product_type:
            warnings.append(f"Filename says {path.stem!r}; workbook Product Type is {product_type!r}")
        return {
            **semantic,
            # Supplemental source evidence does not rewrite the existing v2
            # schema hash or a product's pinned category assignment.
            "sourceMetadata": {
                "version": 1,
                "sourceSha256": source_hash,
                "definitions": documentation,
                "suggestedChoiceKeys": [f["key"] for f in fields if f["column"] in suggested_columns and f.get("allowedValues") and f["attribute"] != "product_type"],
            },
            "definitionCount": len(definitions),
            "definitionRequirementCounts": dict(Counter(d["requirement"] for d in definitions.values())),
            "schemaSha256": schema_hash,
            "sourceSha256": source_hash,
            "sourceFilename": path.name,
            "fieldCount": len(fields),
            "requirementCounts": dict(Counter(field["requirement"] for field in fields)),
            "warnings": warnings,
        }
    finally:
        workbook.close()


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: python3 amazon_template.py path/to/TEMPLATE.xlsm")
    print(json.dumps(extract_template(sys.argv[1]), ensure_ascii=False, separators=(",", ":")))
